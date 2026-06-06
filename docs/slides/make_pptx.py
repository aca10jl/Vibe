#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_pptx.py — 生成《一键一生 × AI 伴读助手 方案》图解 PPT（16:9，可编辑）
依赖：python-pptx>=1.0   运行：python3 make_pptx.py
输出：一键一生-AI助手-方案.pptx
"""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
import os

# ---------- 主题色 ----------
BG     = RGBColor(0x0B, 0x12, 0x20)
PANEL  = RGBColor(0x14, 0x22, 0x38)
PANEL2 = RGBColor(0x0F, 0x18, 0x28)
LINE   = RGBColor(0x24, 0x3A, 0x52)
TEXT   = RGBColor(0xE6, 0xEE, 0xF7)
MUTED  = RGBColor(0x9F, 0xB3, 0xC8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ACC    = RGBColor(0x3B, 0x82, 0xF6)  # 蓝
ACC2   = RGBColor(0x1F, 0x9D, 0x55)  # 绿
WARN   = RGBColor(0xF0, 0xA0, 0x20)  # 橙
BAD    = RGBColor(0xE0, 0x53, 0x3D)  # 红
PUR    = RGBColor(0x7B, 0x2F, 0xF7)  # 紫
FONT   = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width = In(13.333)
prs.slide_height = In(7.5)
W, H = 13.333, 7.5
BLANK = prs.slide_layouts[6]

# ---------- 基础绘图工具 ----------
def slide():
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, BG)
    return s

def rect(s, x, y, w, h, fill, line=None, lw=1.0, rounded=False, radius=0.06):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        In(x), In(y), In(w), In(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    if rounded:
        try: shp.adjustments[0] = radius
        except Exception: pass
    return shp

def shape(s, mso, x, y, w, h, fill, line=None, lw=1.0):
    shp = s.shapes.add_shape(mso, In(x), In(y), In(w), In(h))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    return shp

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = In(0.04); tf.margin_right = In(0.04)
    tf.margin_top = In(0.02); tf.margin_bottom = In(0.02)
    return tf

def para(tf, segs, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=4, space_before=0, bullet=False, lh=None):
    """segs: 字符串 或 [(text,color,bold,size?), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if lh: p.line_spacing = lh
    if isinstance(segs, str):
        segs = [(segs, color, bold, size)]
    for seg in segs:
        t = seg[0]; col = seg[1] if len(seg) > 1 and seg[1] else color
        bd = seg[2] if len(seg) > 2 else bold
        sz = seg[3] if len(seg) > 3 and seg[3] else size
        r = p.add_run(); r.text = ('•  ' + t) if bullet else t
        r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bd; r.font.name = FONT
    return p

def head(s, kicker, title, idx=None):
    rect(s, 0, 0, W, 0.12, ACC)  # 顶部细条
    t = tb(s, 0.7, 0.42, 11.5, 0.4)
    para(t, kicker, size=12, color=ACC, bold=True, first=True)
    t2 = tb(s, 0.7, 0.74, 12.0, 0.7)
    para(t2, title, size=27, color=WHITE, bold=True, first=True)
    rect(s, 0.72, 1.5, 1.15, 0.06, ACC2)
    foot(s, idx)

def foot(s, idx=None):
    f = tb(s, 0.7, 7.04, 8, 0.32)
    para(f, [('◧ 一键一生 × AI 伴读助手', MUTED, False, 10)], first=True)
    if idx is not None:
        r = tb(s, 11.6, 7.04, 1.1, 0.32)
        para(r, [(f'{idx:02d} / 14', MUTED, False, 10)], align=PP_ALIGN.RIGHT, first=True)

def card(s, x, y, w, h, title, lines, accent=ACC, ticon=''):
    rect(s, x, y, w, h, PANEL, line=LINE, lw=1, rounded=True, radius=0.05)
    rect(s, x, y + 0.12, 0.07, h - 0.24, accent)  # 左侧色条
    tf = tb(s, x + 0.22, y + 0.16, w - 0.34, h - 0.3)
    para(tf, [((ticon + ' ' if ticon else '') + title, WHITE, True, 15)], first=True, space_after=6)
    for ln in lines:
        if isinstance(ln, tuple):
            para(tf, ln[0], size=12, color=ln[1] if len(ln) > 1 else MUTED, bullet=True, space_after=3, lh=1.05)
        else:
            para(tf, [(ln, MUTED, False, 12)], bullet=True, space_after=3, lh=1.05)

def fbox(s, x, y, w, h, title, sub='', fill=PANEL, line=ACC, tcol=WHITE, scol=MUTED, ts=13):
    rect(s, x, y, w, h, fill, line=line, lw=1.25, rounded=True, radius=0.08)
    tf = tb(s, x + 0.08, y, w - 0.16, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(title, tcol, True, ts)], align=PP_ALIGN.CENTER, first=True, space_after=2)
    if sub:
        para(tf, [(sub, scol, False, 10)], align=PP_ALIGN.CENTER, space_after=0, lh=1.0)

def arrow(s, x, y, w, h, direction='right', fill=ACC):
    mso = {'right': MSO_SHAPE.RIGHT_ARROW, 'down': MSO_SHAPE.DOWN_ARROW,
           'left': MSO_SHAPE.LEFT_ARROW, 'up': MSO_SHAPE.UP_ARROW}[direction]
    shp = shape(s, mso, x, y, w, h, fill)
    return shp

def table(s, x, y, w, h, data, col_w=None, hfill=ACC, fsz=11, hsz=12, valign=MSO_ANCHOR.MIDDLE):
    rows, cols = len(data), len(data[0])
    g = s.shapes.add_table(rows, cols, In(x), In(y), In(w), In(h)).table
    g.first_row = False; g.horz_banding = False
    if col_w:
        for i, cw in enumerate(col_w):
            g.columns[i].width = In(cw)
    for r in range(rows):
        for c in range(cols):
            cell = g.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = hfill if r == 0 else (PANEL if r % 2 else PANEL2)
            cell.vertical_anchor = valign
            cell.margin_left = In(0.1); cell.margin_right = In(0.06)
            cell.margin_top = In(0.03); cell.margin_bottom = In(0.03)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            val = data[r][c]
            col = WHITE if r == 0 else TEXT
            if val.startswith('✓'): col = ACC2
            elif val.startswith('△'): col = WARN
            rr = p.add_run(); rr.text = val
            rr.font.size = Pt(hsz if r == 0 else fsz); rr.font.bold = (r == 0 or c == 0)
            rr.font.color.rgb = col if not (c == 0 and r > 0) else WHITE
            rr.font.name = FONT
    return g

def chip(s, x, y, w, h, text, fill=PANEL2, col=TEXT, line=LINE):
    rect(s, x, y, w, h, fill, line=line, lw=1, rounded=True, radius=0.5)
    tf = tb(s, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(text, col, True, 11)], align=PP_ALIGN.CENTER, first=True)

# =========================================================
# Slide 1 — 封面
# =========================================================
s = slide()
rect(s, 0, 0, W, H, BG)
# 背景装饰
rect(s, 0, 0, W, 0.16, ACC)
rect(s, 0, H - 0.1, W, 0.1, ACC2)
shape(s, MSO_SHAPE.OVAL, 10.3, -1.6, 4.5, 4.5, PANEL)
shape(s, MSO_SHAPE.OVAL, 11.2, 4.4, 3.6, 3.6, PANEL2)
tf = tb(s, 0.9, 1.7, 9.5, 0.5)
para(tf, [('方案提案 · PROPOSAL', ACC, True, 14)], first=True)
tf = tb(s, 0.9, 2.2, 11.2, 2.0)
para(tf, [('一键一生 ', WHITE, True, 46), ('× ', ACC, True, 46), ('AI 伴读助手', ACC, True, 46)], first=True, space_after=6)
para(tf, [('让晶圆良率根因分析  会读 · 会想 · 会动手', TEXT, False, 22)], space_after=0)
tf = tb(s, 0.9, 4.3, 11.4, 0.6)
para(tf, [('面向半导体 Wafer 全生命周期根因分析的智能体方案 —— 将 AI 助手嵌入现有「一键一生」框架', MUTED, False, 14)], first=True)
for i, (txt, c) in enumerate([('👁 伴读', ACC), ('🔍 分析', ACC2), ('💡 建议', WARN), ('🤖 会动手', PUR), ('🧠 学习习惯', ACC)]):
    chip(s, 0.9 + i * 1.75, 5.3, 1.6, 0.5, txt, fill=PANEL, col=c)
tf = tb(s, 0.9, 6.4, 10, 0.4)
para(tf, [('目标：以低风险、可私有化、即插即用的方式，提升提交正确率、根因定位效率与经验沉淀', MUTED, False, 12)], first=True)

# =========================================================
# Slide 2 — 目录
# =========================================================
s = slide(); head(s, 'AGENDA', '内容提要', 2)
items = [
    ('01', '背景与痛点', '一键一生为何需要 AI 助手'),
    ('02', '方案总览', '三栏布局 + Agent 引擎'),
    ('03', '核心能力', '伴读 / 分析 / 建议 + 会动手'),
    ('04', '技术架构', '前端智能体 + 后端代理'),
    ('05', '集成方式', '低侵入嵌入现有框架'),
    ('06', '兼容场景与优势', '私有化 · 合规 · 即插即用'),
    ('07', '路线图与行动', '灰度落地与风险对策'),
]
x0, y0, cw, ch, gap = 0.8, 1.9, 5.7, 1.45, 0.25
for i, (no, t, d) in enumerate(items):
    col = i % 2; row = i // 2
    x = x0 + col * (cw + 0.4); y = y0 + row * (ch + gap) * 0.78
    if i == 6:
        x = x0 + (cw + 0.4) / 2
    rect(s, x, y, cw, ch * 0.78, PANEL, line=LINE, rounded=True, radius=0.06)
    tf = tb(s, x + 0.2, y, 1.1, ch * 0.78, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(no, ACC, True, 30)], first=True)
    tf = tb(s, x + 1.35, y, cw - 1.5, ch * 0.78, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(t, WHITE, True, 17)], first=True, space_after=2)
    para(tf, [(d, MUTED, False, 12)])

# =========================================================
# Slide 3 — 背景与痛点
# =========================================================
s = slide(); head(s, '01  背景与痛点', '一键一生很强，但"用好它"有门槛', 3)
tf = tb(s, 0.7, 1.7, 12, 0.5)
para(tf, [('现有「一键一生」页面选项丰富（场景/呈现/分析），能一键还原晶圆全链路；但配置链路长、结果信息密度大，"会用、用对、读懂"高度依赖经验。', MUTED, False, 13)], first=True)
pains = [
    ('配置门槛高', ['选项多、必填项易漏', '新手易错填，返工成本高'], BAD, '⚠'),
    ('结果难解读', ['晶圆图/趋势/缺陷/根因并列', '快速判读依赖资深工程师'], WARN, '📉'),
    ('经验依赖强', ['根因判断靠"老师傅"', '知识难以标准化沉淀'], PUR, '🧓'),
    ('操作重复', ['常用配置反复手填', '个人习惯无法被系统记住'], ACC, '🔁'),
]
for i, (t, ls, c, ic) in enumerate(pains):
    card(s, 0.7 + i * 3.05, 2.35, 2.85, 2.05, t, ls, accent=c, ticon=ic)
rect(s, 0.7, 4.75, 11.93, 1.7, PANEL2, line=ACC2, lw=1.25, rounded=True, radius=0.04)
tf = tb(s, 1.0, 4.9, 11.4, 1.45, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [('机会 → 在页面右侧嵌入 AI 助手，恰好补齐这四块短板：', ACC2, True, 15)], first=True, space_after=6)
para(tf, [('伴读', ACC, True, 13), ('  降低解读门槛   ·   ', MUTED, False, 13),
          ('建议', ACC, True, 13), ('  保证正确提交   ·   ', MUTED, False, 13),
          ('分析', ACC, True, 13), ('  沉淀专家经验   ·   ', MUTED, False, 13),
          ('智能体', ACC, True, 13), ('  减少重复操作并持续学习习惯', MUTED, False, 13)])

# =========================================================
# Slide 4 — 方案总览（三栏 + Agent 引擎）
# =========================================================
s = slide(); head(s, '02  方案总览', '一张图看懂：低侵入嵌入页面右侧', 4)
# 三栏
fbox(s, 0.7, 1.85, 3.7, 2.5, '左 · 配置区', '逐步向导\n①场景 ②维度 ③呈现 ④提交', line=ACC2, ts=15)
fbox(s, 4.7, 1.85, 4.3, 2.5, '中 · 结果呈现', '晶圆Bin图 · 良率/WAT趋势\n缺陷Pareto · 工艺时间轴 · 根因候选', line=ACC, ts=15)
fbox(s, 9.3, 1.85, 3.3, 2.5, '右 · AI 助手', '伴读 / 分析 / 建议\n对话 + 习惯画像', line=PUR, ts=15)
# 引擎条
rect(s, 0.7, 4.7, 11.9, 1.55, PANEL, line=LINE, rounded=True, radius=0.04)
tf = tb(s, 0.9, 4.8, 11.5, 0.4)
para(tf, [('Agent 引擎（纯前端，可嵌入任意页面）', WARN, True, 13)], first=True)
eng = [('上下文采集\n读页面状态', ACC), ('工具调用\n真正改页面', ACC2), ('习惯学习\n沉淀画像', PUR), ('三来源 LLM\n规则/本地/Claude', WARN)]
for i, (t, c) in enumerate(eng):
    fbox(s, 0.95 + i * 2.92, 5.25, 2.7, 0.85, t.split('\n')[0], t.split('\n')[1], line=c, ts=12)
# 说明
tf = tb(s, 0.7, 6.35, 12, 0.4)
para(tf, [('嵌入方式：新增一个右侧 aside + 一组脚本，', MUTED, False, 12), ('不改动现有页面的提交与渲染逻辑', ACC2, True, 12), ('；AI 读取现有 DOM/状态作为上下文。', MUTED, False, 12)], first=True)

# =========================================================
# Slide 5 — 核心能力
# =========================================================
s = slide(); head(s, '03  核心能力（一）', '伴读 · 分析 · 建议 —— 把页面"讲明白"', 5)
caps = [
    ('👁 伴读 Companion', ACC, ['读懂当前页面状态与各面板含义', '区分"配置中 / 已出结果"两态', '提交后自动伴读，降低解读门槛'],
     '例：“良率 85.5%，缺陷呈中心聚集，右侧从上到下是…”'),
    ('🔍 分析 Analyze', ACC2, ['基于 commonality 打分做根因分析', '输出按关联度排序的候选 + 置信度', '自动高亮对应结果面板'],
     '例：“最可能根因：CMP-02，置信度 86%”'),
    ('💡 建议 Advise', WARN, ['校验填表完整性并指出问题', '一键修正必填项后再提交', '给出隔离/复测/扩样等处置动作'],
     '例：“缺 Wafer ID 与维度，已帮你补全→可提交”'),
]
for i, (t, c, ls, eg) in enumerate(caps):
    x = 0.7 + i * 4.05
    rect(s, x, 1.85, 3.8, 3.6, PANEL, line=LINE, rounded=True, radius=0.05)
    rect(s, x, 1.85, 3.8, 0.62, c, rounded=True, radius=0.18)
    rect(s, x, 2.2, 3.8, 0.27, PANEL)  # 盖住圆角下半，形成上圆角条
    tf = tb(s, x + 0.2, 1.9, 3.5, 0.5, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(t, WHITE, True, 15)], first=True)
    tf = tb(s, x + 0.25, 2.65, 3.35, 2.0)
    for ln in ls:
        para(tf, [(ln, TEXT, False, 12.5)], bullet=True, space_after=6, lh=1.05)
    rect(s, x + 0.2, 4.75, 3.4, 0.6, PANEL2, line=c, rounded=True, radius=0.12)
    tf = tb(s, x + 0.32, 4.75, 3.2, 0.6, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(eg, MUTED, False, 10.5)], first=True, lh=1.0)
tf = tb(s, 0.7, 5.7, 12, 0.6)
para(tf, [('三种能力共享同一"页面上下文 + 工具"，回答均基于真实数据，不编造数值。', MUTED, False, 13)], first=True)

# =========================================================
# Slide 6 — 智能体：会动手
# =========================================================
s = slide(); head(s, '03  核心能力（二）', '智能体：不只是聊天框，能"动手"操作页面', 6)
# 流程
flow = [('用户/意图', PANEL, MUTED), ('LLM\n页面上下文+工具', ACC, WHITE),
        ('文本 + 工具调用', PANEL, TEXT), ('前端执行\n真正改 DOM', ACC2, WHITE), ('页面更新', PANEL, TEXT)]
bx, by, bw, bh = 0.7, 2.1, 2.05, 1.1
for i, (t, fill, col) in enumerate(flow):
    x = bx + i * (bw + 0.45)
    fbox(s, x, by, bw, bh, t.split('\n')[0], t.split('\n')[1] if '\n' in t else '', fill=fill, line=(ACC if fill == PANEL else fill), tcol=col, ts=12)
    if i < len(flow) - 1:
        arrow(s, x + bw + 0.04, by + bh / 2 - 0.16, 0.38, 0.32, 'right', ACC)
# 回流箭头说明
arrow(s, 6.0, 3.5, 0.32, 0.5, 'up', MUTED)
tf = tb(s, 0.7, 3.95, 12, 0.4)
para(tf, [('循环：可继续追问，AI 依据更新后的页面再决策（单轮可执行，生产可升级为多轮 agentic loop）。', MUTED, False, 12)], first=True)
# 工具清单
rect(s, 0.7, 4.5, 11.93, 1.95, PANEL, line=LINE, rounded=True, radius=0.04)
tf = tb(s, 0.95, 4.62, 11, 0.4)
para(tf, [('🔧 已注册的页面工具（白名单，仅安全动作）', WARN, True, 13)], first=True)
tools = [('set_scenario', '填写/修正场景'), ('set_analysis_dims', '设置分析维度'),
         ('run_lifecycle', '提交一键一生'), ('focus_panel', '定位/高亮面板'), ('apply_my_habits', '套用我的习惯')]
for i, (n, d) in enumerate(tools):
    x = 0.95 + (i % 5) * 2.32
    fbox(s, x, 5.15, 2.18, 1.05, n, d, fill=PANEL2, line=ACC2, ts=11.5)
tf = tb(s, 0.7, 6.55, 12, 0.4)
para(tf, [('兼容层：支持原生 function-calling 的模型直接用；不支持的本地模型解析 ```action {json}``` 动作块 —— 同样能动手。', MUTED, False, 11.5)], first=True)

# =========================================================
# Slide 7 — 持续学习用户习惯
# =========================================================
s = slide(); head(s, '03  核心能力（三）', '智能体：持续学习用户习惯', 7)
loop = [('行为埋点', '配置/提交/聚焦/意图', ACC),
        ('习惯画像', '常用Fab·节点·维度·意图', PUR),
        ('个性化反哺', '默认值 / 建议', ACC2),
        ('用户采纳', '一键套用习惯', WARN)]
cx, cy = 6.6, 3.55
pos = [(3.0, 2.0), (8.4, 2.0), (8.4, 4.6), (3.0, 4.6)]
for i, ((t, d, c), (x, y)) in enumerate(zip(loop, pos)):
    fbox(s, x - 1.35, y - 0.5, 2.7, 1.0, t, d, line=c, ts=13)
# 环形箭头
arrow(s, 5.7, 1.85, 1.6, 0.34, 'right', MUTED)
arrow(s, 8.9, 3.0, 0.34, 1.2, 'down', MUTED)
arrow(s, 5.7, 5.25, 1.6, 0.34, 'left', MUTED)
arrow(s, 3.1, 3.0, 0.34, 1.2, 'up', MUTED)
fbox(s, cx - 1.0, cy - 0.45, 2.0, 0.9, '🧠 越用越懂你', '', fill=PANEL2, line=ACC, ts=13)
# 右侧收益
card(s, 9.7, 2.0, 3.0, 3.4, '带来什么', [
    ('新手 → 资深的经验沉淀', ACC2), ('减少重复手填，提效', ACC2),
    ('个性化默认值更贴合', ACC2), ('画像本地存储，隐私可控', ACC2),
    ('生产可升级为跨会话记忆', MUTED)], accent=PUR, ticon='✨')

# =========================================================
# Slide 8 — 三种 AI 来源
# =========================================================
s = slide(); head(s, '04  技术架构（一）', '三种 AI 来源 · 一键切换 · 私有化友好', 8)
data = [
    ['AI 来源', '部署形态', '数据出域', '工具调用', '适用场景'],
    ['本地推理（规则）', '纯前端，零依赖', '不出域', '✓ 支持', '零配置 Demo / 在线兜底'],
    ['本地大模型', 'Ollama/vLLM/LM Studio/llama.cpp', '内网不出域', '✓ 原生或动作块', '私有化 · 合规 · 离线'],
    ['Claude 在线', '后端代理 claude-opus-4-8', '出域（密钥仅后端）', '✓ 原生 tools', '能力最强 · 复杂分析'],
]
table(s, 0.7, 2.0, 11.93, 2.5, data, col_w=[2.5, 3.5, 2.2, 2.0, 1.73], fsz=11.5, hsz=12.5)
rect(s, 0.7, 4.8, 11.93, 1.5, PANEL2, line=ACC, rounded=True, radius=0.04)
tf = tb(s, 1.0, 4.92, 11.4, 1.3, anchor=MSO_ANCHOR.MIDDLE)
para(tf, [('关键点：', ACC, True, 14), ('三来源统一返回 {text, toolCalls, followups}，对 UI 透明；', TEXT, False, 13)], first=True, space_after=5)
para(tf, [('• 数据敏感/内网场景 → 本地大模型，全程不出域；  • 需最强分析 → Claude；  • 演示/兜底 → 本地推理。', MUTED, False, 12.5)])
para(tf, [('• 在线来源失败自动回退本地推理，可用性有保障。', MUTED, False, 12.5)])

# =========================================================
# Slide 9 — 技术架构图
# =========================================================
s = slide(); head(s, '04  技术架构（二）', '前端智能体 + 后端代理（密钥仅在后端）', 9)
# 浏览器框
rect(s, 0.7, 1.85, 8.3, 4.5, PANEL2, line=LINE, rounded=True, radius=0.03)
tf = tb(s, 0.9, 1.9, 6, 0.35); para(tf, [('浏览器（单页应用，前端零密钥可跑）', MUTED, True, 12)], first=True)
fbox(s, 0.95, 2.35, 2.5, 0.8, '左 · 配置向导', '', line=ACC2, ts=12)
fbox(s, 3.6, 2.35, 2.7, 0.8, '中 · 结果呈现', '', line=ACC, ts=12)
fbox(s, 6.45, 2.35, 2.35, 0.8, '右 · AI 助手', '', line=PUR, ts=12)
rect(s, 0.95, 3.35, 7.85, 1.5, PANEL, line=WARN, rounded=True, radius=0.04)
tf = tb(s, 1.1, 3.42, 7.5, 0.35); para(tf, [('Agent 引擎', WARN, True, 12)], first=True)
for i, t in enumerate(['ContextCollector', 'AgentTools', 'HabitProfile', 'LLMClient']):
    fbox(s, 1.05 + i * 1.92, 3.85, 1.8, 0.85, t, '', fill=PANEL2, line=ACC, ts=11)
fbox(s, 0.95, 5.05, 3.8, 0.7, 'localStorage', '习惯画像持久化', fill=PANEL2, line=PUR, ts=11)
fbox(s, 5.0, 5.05, 3.8, 0.7, 'DOM 状态 ↔ 工具', '读页面 / 改页面', fill=PANEL2, line=ACC2, ts=11)
# 箭头到后端
arrow(s, 9.05, 3.4, 0.7, 0.4, 'right', ACC)
tf = tb(s, 8.95, 2.95, 1.6, 0.4); para(tf, [('/api/chat', MUTED, False, 10)], first=True, align=PP_ALIGN.CENTER)
# 后端框
rect(s, 9.85, 1.85, 2.8, 4.5, PANEL, line=ACC, rounded=True, radius=0.04)
tf = tb(s, 10.0, 1.95, 2.5, 0.4); para(tf, [('Node 后端代理', ACC, True, 12)], first=True, align=PP_ALIGN.CENTER)
fbox(s, 10.05, 2.5, 2.4, 0.75, '静态托管 web/', '', fill=PANEL2, line=LINE, ts=11)
fbox(s, 10.05, 3.4, 2.4, 0.95, 'Provider 路由', '注入领域提示+上下文+工具', fill=PANEL2, line=WARN, ts=11)
fbox(s, 10.05, 4.5, 2.4, 0.75, 'Claude API', 'claude-opus-4-8', fill=PANEL2, line=ACC2, ts=11)
fbox(s, 10.05, 5.4, 2.4, 0.75, '本地大模型', 'OpenAI 兼容端点', fill=PANEL2, line=PUR, ts=11)
tf = tb(s, 0.7, 6.5, 12, 0.4)
para(tf, [('要点：前端默认本地推理即可运行；密钥/端点仅在后端；上下文只传必要派生摘要，避免敏感原始数据外泄。', MUTED, False, 12)], first=True)

# =========================================================
# Slide 10 — 集成方式
# =========================================================
s = slide(); head(s, '05  集成方式', '低侵入接入现有「一键一生」框架（3 步）', 10)
steps = [
    ('① 嵌入右侧栏', ACC, ['新增一个 aside + 一组脚本', '不改现有提交/渲染逻辑', '可灰度：开关控制显隐']),
    ('② 适配上下文契约', ACC2, ['ContextCollector 读现有 DOM/状态', '字段映射到统一快照', 'AI 即可"看懂"你的页面']),
    ('③ 注册页面工具', PUR, ['把现有"设置/提交"动作', '包装为白名单工具', 'AI 即可"动手"操作']),
]
for i, (t, c, ls) in enumerate(steps):
    x = 0.7 + i * 4.05
    card(s, x, 1.9, 3.8, 2.5, t, ls, accent=c)
    if i < 2:
        arrow(s, x + 3.82, 2.95, 0.32, 0.4, 'right', MUTED)
# 数据契约
rect(s, 0.7, 4.7, 11.93, 1.75, PANEL2, line=ACC2, rounded=True, radius=0.03)
tf = tb(s, 0.95, 4.8, 11.4, 0.4); para(tf, [('数据契约（替换合成数据即接真实数据）', ACC2, True, 13)], first=True)
fbox(s, 0.95, 5.25, 3.4, 1.0, '真实 MES/YMS/EDA', '晶圆一生数据', fill=PANEL, line=ACC, ts=12)
arrow(s, 4.45, 5.6, 0.5, 0.34, 'right', MUTED)
fbox(s, 5.05, 5.25, 3.4, 1.0, 'generateLifecycle() 契约', 'meta/summary/waferMap/…', fill=PANEL, line=WARN, ts=12)
arrow(s, 8.55, 5.6, 0.5, 0.34, 'right', MUTED)
fbox(s, 9.15, 5.25, 3.4, 1.0, '可视化 + AI 复用', '无需改动', fill=PANEL, line=ACC2, ts=12)

# =========================================================
# Slide 11 — 兼容场景矩阵
# =========================================================
s = slide(); head(s, '06  兼容场景（一）', '一套方案，覆盖多种部署/用户/任务', 11)
data = [
    ['维度 \\ 场景', '场景 A', '场景 B', '场景 C'],
    ['部署', '公有云（Claude）', '私有化内网（本地大模型）', '完全离线（本地推理）'],
    ['用户', '新手：引导式填表', '资深：秒级根因提效', '混合：习惯自适应'],
    ['分析任务', '单片晶圆追溯', '批次 excursion 排查', '设备 commonality / 参数漂移'],
    ['数据合规', '出域可控（密钥后端）', '✓ 全程不出内网', '✓ 无网络依赖'],
    ['可用性', '在线最强', '在线 + 自动兜底', '✓ 始终可用'],
]
table(s, 0.7, 2.0, 11.93, 3.6, data, col_w=[2.4, 3.18, 3.18, 3.17], fsz=12, hsz=12.5)
tf = tb(s, 0.7, 5.9, 12, 0.6)
para(tf, [('结论：从"零网络离线"到"公有云最强模型"全谱兼容；新手与资深、单片与批次、设备与参数维度均可承载。', MUTED, False, 13)], first=True)

# =========================================================
# Slide 12 — 优势（前后对比）
# =========================================================
s = slide(); head(s, '06  兼容场景（二）', '集成前 vs 集成后：价值一目了然', 12)
data = [
    ['环节', '集成前', '集成后（嵌入 AI 助手）'],
    ['任务提交', '必填项易漏/错填，返工', '✓ 实时校验 + 一键修正再提交'],
    ['结果解读', '靠人工经验逐图判读', '✓ AI 伴读，逐面板讲解'],
    ['根因定位', '依赖资深工程师', '✓ 候选 + 置信度，秒级给出'],
    ['经验沉淀', '存于个人脑中', '✓ 习惯画像，可复用/可成长'],
    ['上手时间', '长，培训成本高', '✓ 显著缩短，新手即用'],
    ['数据合规', '—', '✓ 私有化 / 离线可选'],
]
table(s, 0.7, 2.0, 11.93, 4.1, data, col_w=[2.3, 4.6, 5.03], fsz=12.5, hsz=13)
tf = tb(s, 0.7, 6.35, 12, 0.4)
para(tf, [('提交正确率 ↑   ·   解读门槛 ↓   ·   根因定位提速   ·   经验资产化   ·   合规可控', ACC2, True, 13.5)], first=True, align=PP_ALIGN.CENTER)

# =========================================================
# Slide 13 — 路线图 + 风险
# =========================================================
s = slide(); head(s, '07  路线图与风险', '低风险灰度落地，分阶段交付', 13)
phases = [
    ('P0 · 1–2 周', ACC, ['嵌入右侧栏 + 本地推理', '对接数据契约', '零风险灰度']),
    ('P1 · 2–4 周', ACC2, ['接本地大模型 / Claude', '工具适配现有动作', '试点产品线验证']),
    ('P2 · 进阶', WARN, ['流式逐字输出', '跨会话记忆(Memory)', '体验升级']),
    ('P3 · 增强', PUR, ['RAG：SOP/历史 FA', '采纳信号在线学习', '越用越准']),
]
for i, (t, c, ls) in enumerate(phases):
    x = 0.7 + i * 3.05
    rect(s, x, 1.95, 2.85, 0.5, c, rounded=True, radius=0.2)
    tf = tb(s, x, 1.95, 2.85, 0.5, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(t, WHITE, True, 13)], align=PP_ALIGN.CENTER, first=True)
    card(s, x, 2.55, 2.85, 1.95, '', ls, accent=c)
    if i < 3:
        arrow(s, x + 2.87, 2.05, 0.3, 0.3, 'right', MUTED)
# 风险对策
rect(s, 0.7, 4.85, 11.93, 1.6, PANEL2, line=BAD, rounded=True, radius=0.03)
tf = tb(s, 0.95, 4.95, 11.4, 0.4); para(tf, [('风险与对策', BAD, True, 13)], first=True)
risks = [('数据安全', '私有化/离线部署，全程不出内网'),
         ('模型幻觉', '基于页面上下文作答+不编造+工具白名单'),
         ('侵入风险', '低耦合嵌入，开关灰度，可随时回退')]
for i, (r, m) in enumerate(risks):
    x = 0.95 + i * 3.92
    fbox(s, x, 5.4, 3.75, 0.9, r, m, fill=PANEL, line=WARN, ts=12)

# =========================================================
# Slide 14 — 行动号召
# =========================================================
s = slide()
rect(s, 0, 0, W, H, BG); rect(s, 0, 0, W, 0.16, ACC); rect(s, 0, H - 0.1, W, 0.1, ACC2)
shape(s, MSO_SHAPE.OVAL, 10.0, 4.0, 4.4, 4.4, PANEL2)
tf = tb(s, 0.9, 1.4, 11, 0.5); para(tf, [('行动号召 · CALL TO ACTION', ACC, True, 14)], first=True)
tf = tb(s, 0.9, 1.95, 11.4, 1.4)
para(tf, [('低风险 · 可私有化 · 即插即用 · 可成长', WHITE, True, 32)], first=True, space_after=6)
para(tf, [('建议将 AI 伴读助手作为「一键一生」的标准能力集成', TEXT, False, 18)])
card(s, 0.9, 3.7, 5.6, 2.7, '我们请求', [
    ('批准 P0 灰度接入（约 2 周）', ACC2),
    ('指定 1 个试点产品线 / 页面', ACC2),
    ('明确数据合规边界（私有化/离线）', ACC2),
    ('指派 1 名前端 + 1 名数据对接', ACC2)], accent=ACC, ticon='✅')
card(s, 6.8, 3.7, 5.6, 2.7, '你将获得', [
    ('提交正确率与解读效率提升', WARN),
    ('根因定位提速、经验资产化', WARN),
    ('已可运行的 Demo（本地零配置）', WARN),
    ('完整源码 + 调研设计 + 部署文档', WARN)], accent=ACC2, ticon='🎁')
tf = tb(s, 0.9, 6.55, 11.5, 0.5)
para(tf, [('Demo 与文档：仓库 web/（一键运行）· docs/（调研与设计、部署与Demo）', MUTED, False, 12)], first=True)

# ---------- 保存 ----------
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '一键一生-AI助手-方案.pptx')
prs.save(out)
print('Saved:', out, '| slides:', len(prs.slides._sldIdLst))
