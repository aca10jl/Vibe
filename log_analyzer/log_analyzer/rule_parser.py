"""
规则解析器：从 Excel 文件中加载错误匹配规则。

支持的 Excel 格式：
- 必选列: "error match" (正则表达式)
- 可选列: "error description" (错误描述), "severity" (严重程度), "category" (分类)
"""

import re
import json
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Optional


@dataclass
class ErrorRule:
    """一条错误匹配规则"""
    pattern_str: str
    pattern: re.Pattern = field(repr=False, compare=False, default=None)
    description: str = ""
    severity: str = "ERROR"
    category: str = "general"
    rule_id: int = 0

    def __post_init__(self):
        if self.pattern is None:
            self.pattern = re.compile(self.pattern_str, re.IGNORECASE)

    def to_dict(self) -> dict:
        d = asdict(self)
        d.pop("pattern")
        return d


def load_rules_from_excel(excel_path: str) -> list[ErrorRule]:
    """从 Excel 文件加载规则"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("请安装 openpyxl: pip install openpyxl")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active

    # 读取表头，找到列索引
    headers = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=False))):
        val = str(cell.value).strip().lower() if cell.value else ""
        headers[val] = idx

    # 查找关键列
    match_col = None
    desc_col = None
    severity_col = None
    category_col = None

    for name, idx in headers.items():
        if "error match" in name or "match" in name or "regex" in name or "pattern" in name:
            match_col = idx
        if "error description" in name or "description" in name or "desc" in name:
            desc_col = idx
        if "severity" in name or "level" in name:
            severity_col = idx
        if "category" in name or "type" in name or "分类" in name:
            category_col = idx

    if match_col is None:
        raise ValueError(
            f"Excel 中未找到 'error match' 列。"
            f"可用列: {list(headers.keys())}"
        )

    rules = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        pattern_str = row[match_col] if match_col < len(row) else None
        if not pattern_str or not str(pattern_str).strip():
            continue

        pattern_str = str(pattern_str).strip()

        # 尝试编译正则，失败则转义为字面量
        try:
            re.compile(pattern_str)
        except re.error:
            pattern_str = re.escape(pattern_str)

        desc = ""
        if desc_col is not None and desc_col < len(row) and row[desc_col]:
            desc = str(row[desc_col]).strip()

        severity = "ERROR"
        if severity_col is not None and severity_col < len(row) and row[severity_col]:
            severity = str(row[severity_col]).strip().upper()

        category = "general"
        if category_col is not None and category_col < len(row) and row[category_col]:
            category = str(row[category_col]).strip()

        rules.append(ErrorRule(
            pattern_str=pattern_str,
            description=desc,
            severity=severity,
            category=category,
            rule_id=row_idx,
        ))

    wb.close()
    return rules


def load_rules_from_json(json_path: str) -> list[ErrorRule]:
    """从 JSON 文件加载规则（备选方案）"""
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    rules = []
    for idx, item in enumerate(data, start=1):
        rules.append(ErrorRule(
            pattern_str=item["pattern"],
            description=item.get("description", ""),
            severity=item.get("severity", "ERROR"),
            category=item.get("category", "general"),
            rule_id=idx,
        ))
    return rules


# 内置的通用错误模式（举一反三）
BUILTIN_PATTERNS = [
    ErrorRule(
        pattern_str=r"(?i)\b(FATAL|PANIC|CRITICAL)\b",
        description="致命级别错误",
        severity="FATAL",
        category="severity",
        rule_id=9001,
    ),
    ErrorRule(
        pattern_str=r"(?i)out\s*of\s*memory|OOM|oom-killer",
        description="内存溢出",
        severity="FATAL",
        category="resource",
        rule_id=9002,
    ),
    ErrorRule(
        pattern_str=r"(?i)segmentation\s+fault|segfault|SIGSEGV",
        description="段错误 / 内存访问违规",
        severity="FATAL",
        category="crash",
        rule_id=9003,
    ),
    ErrorRule(
        pattern_str=r"(?i)stack\s*overflow",
        description="栈溢出",
        severity="FATAL",
        category="crash",
        rule_id=9004,
    ),
    ErrorRule(
        pattern_str=r"(?i)deadlock\s+detected|deadlock",
        description="死锁检测",
        severity="ERROR",
        category="concurrency",
        rule_id=9005,
    ),
    ErrorRule(
        pattern_str=r"(?i)connection\s+(refused|reset|timed?\s*out)|ECONNREFUSED|ECONNRESET|ETIMEDOUT",
        description="连接异常（拒绝/重置/超时）",
        severity="ERROR",
        category="network",
        rule_id=9006,
    ),
    ErrorRule(
        pattern_str=r"(?i)disk\s+(full|space)|no\s+space\s+left",
        description="磁盘空间不足",
        severity="ERROR",
        category="resource",
        rule_id=9007,
    ),
    ErrorRule(
        pattern_str=r"(?i)permission\s+denied|access\s+denied|EACCES",
        description="权限被拒绝",
        severity="ERROR",
        category="permission",
        rule_id=9008,
    ),
    ErrorRule(
        pattern_str=r"(?i)timeout|timed?\s*out",
        description="操作超时",
        severity="WARN",
        category="performance",
        rule_id=9009,
    ),
    ErrorRule(
        pattern_str=r"(?i)(null\s*pointer|NullPointerException|nullptr|nil\s+pointer)",
        description="空指针异常",
        severity="ERROR",
        category="crash",
        rule_id=9010,
    ),
    ErrorRule(
        pattern_str=r"(?i)file\s+not\s+found|FileNotFoundException|ENOENT",
        description="文件未找到",
        severity="ERROR",
        category="io",
        rule_id=9011,
    ),
    ErrorRule(
        pattern_str=r"(?i)too\s+many\s+open\s+files|EMFILE",
        description="打开文件数超限",
        severity="ERROR",
        category="resource",
        rule_id=9012,
    ),
    ErrorRule(
        pattern_str=r"(?i)core\s+dump(ed)?",
        description="产生 core dump",
        severity="FATAL",
        category="crash",
        rule_id=9013,
    ),
    ErrorRule(
        pattern_str=r"(?i)(?:Exception|Error|Traceback)\s*[:(\[]",
        description="通用异常/错误标记",
        severity="ERROR",
        category="exception",
        rule_id=9014,
    ),
]
